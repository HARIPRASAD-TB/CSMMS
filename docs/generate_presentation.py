"""
Generate the MCA Major Project PowerPoint for the
"Construction Service & Material Management System" project.

Widescreen 16:9 deck with a dark construction theme (matches the app UI),
native shape-based diagrams (architecture, roles, ER, workflow), styled
tables, and real application screenshots from docs/screenshots/.

Usage:
    python -m pip install --user python-pptx
    python generate_presentation.py
"""

from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from PIL import Image
except ImportError:
    print("Install dependencies: python -m pip install --user python-pptx")
    raise

HERE = Path(__file__).parent
OUTPUT = HERE / "CSMMS_MCA_Presentation.pptx"
SHOTS = HERE / "screenshots"

# Brand palette (matches BuildConnect UI)
ACCENT = RGBColor(0xFA, 0xCC, 0x15)   # yellow
DARK = RGBColor(0x14, 0x14, 0x14)     # page background
CARD = RGBColor(0x23, 0x23, 0x23)     # card background
CARD2 = RGBColor(0x2E, 0x2E, 0x2E)
WHITE = RGBColor(0xF5, 0xF5, 0xF5)
MUTED = RGBColor(0xA1, 0xA1, 0xAA)
GREEN = RGBColor(0x4A, 0xDE, 0x80)
BLUE = RGBColor(0x60, 0xA5, 0xFA)
ORANGE = RGBColor(0xFB, 0x92, 0x3C)
PINK = RGBColor(0xF4, 0x72, 0xB6)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

FONT = "Segoe UI"
FONT_BOLD = "Segoe UI Semibold"


# ---------------------------------------------------------------- helpers

def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = DARK
    return slide


def no_line(shape):
    shape.line.fill.background()


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    no_line(shape)


def set_text(shape, text, size=14, color=WHITE, bold=False, align=PP_ALIGN.CENTER,
             font=FONT, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run() if not p.runs else p.runs[0]
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def add_box(slide, x, y, w, h, color=CARD, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE,
            radius=0.12):
    shp = slide.shapes.add_shape(shape_type, x, y, w, h)
    solid(shp, color)
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.shadow.inherit = False
    return shp


def add_text_box(slide, x, y, w, h, text, size=14, color=WHITE, bold=False,
                 align=PP_ALIGN.LEFT, font=FONT, line_spacing=1.0, space_after=4):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_arrow(slide, x, y, w, h, color=ACCENT, direction="down"):
    shape_map = {
        "down": MSO_SHAPE.DOWN_ARROW,
        "right": MSO_SHAPE.RIGHT_ARROW,
        "left_right": MSO_SHAPE.LEFT_RIGHT_ARROW,
        "up_down": MSO_SHAPE.UP_DOWN_ARROW,
    }
    shp = slide.shapes.add_shape(shape_map[direction], x, y, w, h)
    solid(shp, color)
    shp.shadow.inherit = False
    return shp


def add_line(slide, x1, y1, x2, y2, color=MUTED, weight=1.5):
    conn = slide.shapes.add_connector(1, x1, y1, x2, y2)  # 1 = straight
    conn.line.color.rgb = color
    conn.line.width = Pt(weight)
    return conn


def header(slide, title, subtitle=None, number=None):
    """Accent bar + slide title + optional kicker, footer strip."""
    add_box(slide, Inches(0.55), Inches(0.42), Inches(0.09), Inches(0.62),
            color=ACCENT, shape_type=MSO_SHAPE.RECTANGLE)
    add_text_box(slide, Inches(0.8), Inches(0.32), Inches(11.2), Inches(0.7),
                 title, size=30, color=WHITE, bold=True, font=FONT_BOLD)
    if subtitle:
        add_text_box(slide, Inches(0.82), Inches(0.98), Inches(11.2), Inches(0.4),
                     subtitle, size=13, color=MUTED)
    # footer
    add_text_box(slide, Inches(0.55), Inches(7.05), Inches(7), Inches(0.35),
                 "Construction Service & Material Management System  |  MCA Major Project",
                 size=9, color=MUTED)
    if number is not None:
        add_text_box(slide, Inches(12.3), Inches(7.05), Inches(0.8), Inches(0.35),
                     str(number), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def bullet_card(slide, x, y, w, h, title, lines, title_color=ACCENT, body_size=13):
    card = add_box(slide, x, y, w, h)
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Pt(14)
    tf.margin_right = Pt(14)
    tf.margin_top = Pt(12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = FONT_BOLD
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = title_color
    p.space_after = Pt(8)
    for line in lines:
        p = tf.add_paragraph()
        p.space_after = Pt(5)
        run = p.add_run()
        run.text = "•  " + line
        run.font.name = FONT
        run.font.size = Pt(body_size)
        run.font.color.rgb = WHITE
    return card


def styled_table(slide, x, y, w, h, headers, rows, col_widths=None,
                 header_size=13, body_size=12):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), x, y, w, h)
    table = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Emu(int(w * cw / total))
    for j, head in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = head
        run.font.name = FONT_BOLD
        run.font.size = Pt(header_size)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x14, 0x14, 0x14)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD if i % 2 == 0 else CARD2
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = val
            run.font.name = FONT
            run.font.size = Pt(body_size)
            run.font.color.rgb = WHITE
    return table


def add_screenshot(slide, img_path, x, y, w, caption=None):
    """Insert image with an accent border frame."""
    img = Image.open(img_path)
    iw, ih = img.size
    h = Emu(int(w * ih / iw))
    pad = Inches(0.04)
    add_box(slide, x - pad, y - pad, w + pad * 2, Emu(int(h) + int(pad) * 2),
            color=ACCENT, radius=0.03)
    slide.shapes.add_picture(str(img_path), x, y, width=w)
    if caption:
        add_text_box(slide, x, y + h + Inches(0.08), w, Inches(0.3),
                     caption, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    return h


def crop_screenshots():
    """Trim the bottom strip (Next.js dev badge) from captured screenshots."""
    if not SHOTS.exists():
        return
    for png in SHOTS.glob("*.png"):
        cropped = SHOTS / "cropped" / png.name
        cropped.parent.mkdir(exist_ok=True)
        img = Image.open(png)
        w, h = img.size
        img.crop((0, 0, w, h - 80)).save(cropped)


def shot(name):
    p = SHOTS / "cropped" / f"{name}.png"
    return p if p.exists() else None


# ---------------------------------------------------------------- slides

def slide_title(prs):
    s = blank_slide(prs)
    # decorative corner shapes
    tri = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(-1.2), Inches(5.2),
                             Inches(4.5), Inches(2.4))
    solid(tri, CARD)
    tri.rotation = 0
    bar = add_box(s, Inches(0), Inches(0), SLIDE_W, Inches(0.18),
                  color=ACCENT, shape_type=MSO_SHAPE.RECTANGLE)
    add_box(s, Inches(0), Inches(7.32), SLIDE_W, Inches(0.18),
            color=ACCENT, shape_type=MSO_SHAPE.RECTANGLE)
    # icon: yellow rounded square badge
    icon = add_box(s, Inches(5.92), Inches(0.8), Inches(1.5), Inches(1.5),
                   color=ACCENT, radius=0.28)
    set_text(icon, "MCA", size=30, color=DARK, bold=True, font=FONT_BOLD)

    add_text_box(s, Inches(0.7), Inches(2.5), Inches(11.93), Inches(1.7),
                 "CONSTRUCTION SERVICE & MATERIAL\nMANAGEMENT SYSTEM",
                 size=38, color=ACCENT, bold=True,
                 align=PP_ALIGN.CENTER, font=FONT_BOLD)
    add_text_box(s, Inches(1.2), Inches(4.2), Inches(10.93), Inches(0.5),
                 "A Unified Web Marketplace for Workers, Contractors & Building Materials",
                 size=16, color=WHITE, align=PP_ALIGN.CENTER)

    card = add_box(s, Inches(3.42), Inches(4.85), Inches(6.5), Inches(1.95))
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    entries = [
        ("HARIPRASAD T B", 18, WHITE, True),
        ("Enrolment No.: 232MCAN0168   |   Reg. No.: 23MCA5168", 13, MUTED, False),
        ("Guide: Dr. S. Meera  —  Associate Professor & Head (I/c)", 13, WHITE, False),
        ("Department of Computer Science (AI)", 12, MUTED, False),
        ("Centre for Distance and Online Education, Bharathiar University  |  2023–2025",
         12, ACCENT, False),
    ]
    for i, (txt, size, color, bold) in enumerate(entries):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = txt
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def slide_agenda(prs):
    s = blank_slide(prs)
    header(s, "Agenda", number=2)
    items = [
        "Introduction & Problem Statement", "Objectives",
        "Existing vs Proposed System", "System Architecture",
        "User Roles & Modules", "Database Design (ER)",
        "Technology Stack", "Application Walkthrough",
        "Testing & Results", "Conclusion & Future Work",
    ]
    cols, x0, y0 = 2, Inches(1.1), Inches(1.6)
    cw, ch, gx, gy = Inches(5.5), Inches(0.85), Inches(0.55), Inches(0.22)
    for i, item in enumerate(items):
        r, c = divmod(i, cols)
        x = x0 + c * (cw + gx)
        y = y0 + r * (ch + gy)
        chip = add_box(s, x, y, Inches(0.85), ch, color=ACCENT, radius=0.25)
        set_text(chip, f"{i+1:02d}", size=20, color=DARK, bold=True, font=FONT_BOLD)
        card = add_box(s, x + Inches(0.95), y, cw - Inches(0.95), ch)
        set_text(card, item, size=15, color=WHITE, align=PP_ALIGN.LEFT)
        card.text_frame.margin_left = Pt(14)


def slide_intro(prs):
    s = blank_slide(prs)
    header(s, "Introduction", number=3)
    bullet_card(s, Inches(0.7), Inches(1.5), Inches(6.4), Inches(4.9),
                "The Context", [
            "Construction is one of India's largest employment sectors",
            "Hiring relies on word-of-mouth and local referrals",
            "Material purchase needs visits to multiple suppliers",
            "No unified digital platform for services + materials",
            "Providers lack online presence and booking tools",
        ], body_size=15)
    card = add_box(s, Inches(7.4), Inches(1.5), Inches(5.2), Inches(4.9))
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (txt, size, color, bold) in enumerate([
        ("Proposed System", 24, ACCENT, True),
        ("", 8, WHITE, False),
        ("A full-stack web marketplace that unifies", 15, WHITE, False),
        ("worker hiring  •  contractor booking  •  material purchase", 15, ACCENT, True),
        ("", 8, WHITE, False),
        ("Built with Next.js 16, React 19, MongoDB,", 14, MUTED, False),
        ("and JWT-secured role-based access", 14, MUTED, False),
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = txt
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def slide_problem(prs):
    s = blank_slide(prs)
    header(s, "Problem Statement", number=4)
    problems = [
        ("No Verification", "Hard to find verified, rated workers and contractors", PINK),
        ("Opaque Pricing", "Labour and material rates unclear until work begins", ORANGE),
        ("No Records", "Bookings, orders and payments are untracked", BLUE),
        ("No Digital Tools", "Providers cannot manage requests or earnings online", GREEN),
        ("No Oversight", "No central admin control over users and listings", ACCENT),
    ]
    x0, y = Inches(0.7), Inches(2.0)
    cw, gap = Inches(2.34), Inches(0.12)
    for i, (title, body, color) in enumerate(problems):
        x = x0 + i * (cw + gap)
        dot = add_box(s, x + cw / 2 - Inches(0.3), y, Inches(0.6), Inches(0.6),
                      color=color, shape_type=MSO_SHAPE.OVAL)
        set_text(dot, str(i + 1), size=20, color=DARK, bold=True)
        card = add_box(s, x, y + Inches(0.85), cw, Inches(3.0))
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_top = Pt(14)
        tf.margin_left = Pt(10)
        tf.margin_right = Pt(10)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        run.font.name = FONT_BOLD
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = color
        p.space_after = Pt(8)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = body
        run2.font.name = FONT
        run2.font.size = Pt(13)
        run2.font.color.rgb = WHITE


def slide_objectives(prs):
    s = blank_slide(prs)
    header(s, "Objectives", number=5)
    objs = [
        "Develop a unified construction marketplace web application",
        "Implement secure JWT authentication & role-based access (5 roles)",
        "Enable worker hiring and contractor booking with filters & ratings",
        "Provide material e-commerce with cart and checkout",
        "Support bookings, orders, reviews and unified activity history",
        "Deliver Admin, Provider and Vendor management dashboards",
    ]
    y0 = Inches(1.55)
    for i, obj in enumerate(objs):
        y = y0 + i * Inches(0.92)
        chip = add_box(s, Inches(1.0), y, Inches(0.75), Inches(0.75),
                       color=ACCENT, radius=0.3)
        set_text(chip, str(i + 1), size=22, color=DARK, bold=True, font=FONT_BOLD)
        card = add_box(s, Inches(1.95), y, Inches(10.4), Inches(0.75))
        set_text(card, obj, size=16, color=WHITE, align=PP_ALIGN.LEFT)
        card.text_frame.margin_left = Pt(16)


def slide_existing_proposed(prs):
    s = blank_slide(prs)
    header(s, "Existing System vs Proposed System", number=6)
    bullet_card(s, Inches(0.7), Inches(1.55), Inches(5.9), Inches(4.9),
                "EXISTING — Manual & Fragmented", [
            "Word-of-mouth referrals, phone negotiations",
            "Visits to multiple shops for material prices",
            "No ratings, verification or comparison",
            "Paper-based or zero record keeping",
            "Partial apps: services OR materials, never both",
        ], title_color=PINK, body_size=14)
    bullet_card(s, Inches(6.85), Inches(1.55), Inches(5.9), Inches(4.9),
                "PROPOSED — Unified Web System", [
            "Single platform: workers + contractors + materials",
            "Transparent pricing, photos, ratings & reviews",
            "Cart, checkout, bookings and order tracking",
            "Role-based dashboards for every participant",
            "Secure JWT sessions, MongoDB persistence",
        ], title_color=GREEN, body_size=14)
    arrow = add_arrow(s, Inches(6.18), Inches(3.6), Inches(1.0), Inches(0.6),
                      direction="right")
    arrow.fill.fore_color.rgb = ACCENT


def slide_architecture(prs):
    s = blank_slide(prs)
    header(s, "System Architecture", "Three-tier architecture", number=7)
    tiers = [
        ("PRESENTATION LAYER", "Next.js Pages  •  React 19 Components  •  Tailwind CSS  •  Zustand Cart", BLUE),
        ("APPLICATION LAYER", "Next.js API Routes (REST)  •  Middleware  •  JWT Auth  •  Role Permissions", ACCENT),
        ("DATA LAYER", "MongoDB  •  Mongoose ODM  •  users · providers · products · bookings · orders · reviews", GREEN),
    ]
    x, w = Inches(2.2), Inches(8.9)
    y0, th, gap = Inches(1.7), Inches(1.15), Inches(0.5)
    labels = ["HTTP / REST  +  JWT Cookie", "Mongoose Queries"]
    for i, (title, body, color) in enumerate(tiers):
        y = y0 + i * (th + gap)
        card = add_box(s, x, y, w, th, color=CARD)
        card.line.color.rgb = color
        card.line.width = Pt(2)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        run.font.name = FONT_BOLD
        run.font.size = Pt(17)
        run.font.bold = True
        run.font.color.rgb = color
        p.space_after = Pt(4)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = body
        run2.font.name = FONT
        run2.font.size = Pt(13)
        run2.font.color.rgb = WHITE
        if i < 2:
            ay = y + th + Inches(0.02)
            add_arrow(s, x + w / 2 - Inches(0.22), ay, Inches(0.44), gap - Inches(0.04),
                      direction="up_down")
            add_text_box(s, x + w / 2 + Inches(0.4), ay + Inches(0.05),
                         Inches(3.2), Inches(0.4), labels[i], size=11, color=MUTED)
    # side actors
    actor = add_box(s, Inches(0.55), Inches(1.85), Inches(1.35), Inches(0.85),
                    color=CARD2)
    set_text(actor, "Browser\n(All Roles)", size=11, color=MUTED)
    add_arrow(s, Inches(1.92), Inches(2.12), Inches(0.26), Inches(0.3),
              direction="right")


def slide_roles(prs):
    s = blank_slide(prs)
    header(s, "User Roles", "Five roles with dedicated dashboards", number=8)
    hub = add_box(s, Inches(5.42), Inches(3.15), Inches(2.5), Inches(1.3),
                  color=ACCENT, radius=0.2)
    set_text(hub, "Unified\nPlatform", size=17, color=DARK, bold=True,
             font=FONT_BOLD)
    roles = [
        ("CUSTOMER", "Browse • Book • Order\nReview • Track activity", Inches(0.75), Inches(1.55), BLUE),
        ("WORKER", "Listing • Bookings\nEarnings dashboard", Inches(9.75), Inches(1.55), GREEN),
        ("CONTRACTOR", "Projects • Bookings\nEarnings dashboard", Inches(0.75), Inches(4.7), ORANGE),
        ("VENDOR", "Products • Orders\nSales statistics", Inches(9.75), Inches(4.7), PINK),
        ("ADMIN", "Users • Listings • Bookings • Orders — full control", Inches(4.47), Inches(5.55), ACCENT),
    ]
    centers = []
    for title, body, x, y, color in roles:
        w = Inches(2.9) if title != "ADMIN" else Inches(4.4)
        h = Inches(1.45) if title != "ADMIN" else Inches(1.1)
        card = add_box(s, x, y, w, h, color=CARD)
        card.line.color.rgb = color
        card.line.width = Pt(2)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        run.font.name = FONT_BOLD
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = color
        p.space_after = Pt(3)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = body
        run2.font.name = FONT
        run2.font.size = Pt(11.5)
        run2.font.color.rgb = WHITE
        centers.append((x + w / 2, y + h / 2))
    hub_cx, hub_cy = Inches(6.67), Inches(3.8)
    for cx, cy in centers:
        add_line(s, hub_cx, hub_cy, cx, cy, color=RGBColor(0x52, 0x52, 0x52), weight=1.25)
    # redraw hub on top of lines
    hub2 = add_box(s, Inches(5.42), Inches(3.15), Inches(2.5), Inches(1.3),
                   color=ACCENT, radius=0.2)
    set_text(hub2, "Unified\nPlatform", size=17, color=DARK, bold=True,
             font=FONT_BOLD)


def slide_modules(prs):
    s = blank_slide(prs)
    header(s, "Application Modules", number=9)
    modules = [
        ("Authentication", "Register, login, JWT cookie sessions, middleware protection"),
        ("Services", "Workers & contractors with filters, detail pages, booking"),
        ("Materials", "Product catalog, categories, sorting, image sliders"),
        ("Cart & Checkout", "Per-user cart for products + services, order creation"),
        ("My Activity", "Unified history of orders and bookings"),
        ("Reviews & Ratings", "Star ratings and comments for providers"),
        ("Admin Panel", "Users, providers, vendors, materials, bookings, orders"),
        ("Provider Panel", "Booking management and earnings analytics"),
        ("Vendor Dashboard", "Product listings and sales overview"),
    ]
    cols, x0, y0 = 3, Inches(0.7), Inches(1.55)
    cw, ch = Inches(4.0), Inches(1.55)
    gx, gy = Inches(0.27), Inches(0.22)
    colors = [BLUE, GREEN, ORANGE, PINK, ACCENT, BLUE, GREEN, ORANGE, PINK]
    for i, (title, body) in enumerate(modules):
        r, c = divmod(i, cols)
        x = x0 + c * (cw + gx)
        y = y0 + r * (ch + gy)
        card = add_box(s, x, y, cw, ch)
        bar = add_box(s, x, y, Inches(0.1), ch, color=colors[i],
                      shape_type=MSO_SHAPE.RECTANGLE)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Pt(18)
        tf.margin_right = Pt(10)
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.name = FONT_BOLD
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = colors[i]
        p.space_after = Pt(4)
        p2 = tf.add_paragraph()
        run2 = p2.add_run()
        run2.text = body
        run2.font.name = FONT
        run2.font.size = Pt(11.5)
        run2.font.color.rgb = MUTED


def slide_er(prs):
    s = blank_slide(prs)
    header(s, "Database Design — ER Diagram",
           "MongoDB `buildconnect` — six collections", number=10)

    def entity(x, y, name, fields, color):
        w, h = Inches(2.7), Inches(1.7)
        card = add_box(s, x, y, w, h, color=CARD)
        card.line.color.rgb = color
        card.line.width = Pt(2)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_top = Pt(8)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = name
        run.font.name = FONT_BOLD
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = color
        p.space_after = Pt(4)
        for f in fields:
            pf = tf.add_paragraph()
            pf.alignment = PP_ALIGN.CENTER
            rf = pf.add_run()
            rf.text = f
            rf.font.name = FONT
            rf.font.size = Pt(10.5)
            rf.font.color.rgb = MUTED
        return (x + w / 2, y + h / 2, x, y, w, h)

    user = entity(Inches(0.9), Inches(2.7), "USER",
                  ["name, email, mobile", "password (bcrypt)", "role, isBlocked"], ACCENT)
    provider = entity(Inches(5.3), Inches(1.45), "PROVIDER",
                      ["type: worker/contractor", "pricing, portfolio", "rating, isApproved"], GREEN)
    review = entity(Inches(9.7), Inches(1.45), "REVIEW",
                    ["rating (1–5)", "comment", "userId, providerId"], PINK)
    booking = entity(Inches(5.3), Inches(4.3), "BOOKING",
                     ["serviceType, startDate", "totalAmount", "status workflow"], BLUE)
    order = entity(Inches(9.7), Inches(4.3), "ORDER",
                   ["items[], total", "deliveryAddress", "status workflow"], ORANGE)
    product = entity(Inches(9.7), Inches(2.85), "PRODUCT",
                     ["name, category, price", "unit, stock, image", "vendorId"], WHITE)

    pairs = [
        (user, provider, "1 : N  owns"),
        (provider, review, "1 : N  has"),
        (user, booking, "1 : N  places"),
        (user, order, None),
        (provider, booking, None),
        (product, order, "in items[]"),
    ]
    for a, b, label in pairs:
        add_line(s, Emu(int(a[0])), Emu(int(a[1])), Emu(int(b[0])), Emu(int(b[1])),
                 color=RGBColor(0x5A, 0x5A, 0x5A), weight=1.5)
        if label:
            mx = (a[0] + b[0]) / 2
            my = (a[1] + b[1]) / 2
            add_text_box(s, Emu(int(mx - Inches(0.8))), Emu(int(my - Inches(0.32))),
                         Inches(1.6), Inches(0.3), label, size=10, color=ACCENT,
                         align=PP_ALIGN.CENTER)


def slide_collections(prs):
    s = blank_slide(prs)
    header(s, "Database Collections", number=11)
    styled_table(
        s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.8),
        ["Collection", "Purpose", "Key Fields"],
        [
            ["users", "Accounts & credentials (5 roles)", "name, email, password, role, isBlocked"],
            ["providers", "Worker / contractor listings", "type, title, location, pricing, rating, portfolio"],
            ["products", "Material catalog", "name, category, price, unit, stock, vendorId"],
            ["bookings", "Service bookings", "userId, providerId, startDate, totalAmount, status"],
            ["orders", "Material purchases", "userId, items[], total, deliveryAddress, status"],
            ["reviews", "Provider feedback", "userId, providerId, rating (1–5), comment"],
        ],
        col_widths=[2, 3.5, 5],
        header_size=14, body_size=13,
    )


def slide_tech(prs):
    s = blank_slide(prs)
    header(s, "Technology Stack", number=12)
    techs = [
        ("Next.js 16", "Full-stack framework — pages + REST API routes", BLUE),
        ("React 19", "Component-based responsive user interface", BLUE),
        ("Tailwind CSS 4", "Utility-first styling, dark construction theme", GREEN),
        ("MongoDB + Mongoose 8", "Document database with schema validation", GREEN),
        ("JWT + bcrypt", "HTTP-only cookie sessions, hashed passwords", ORANGE),
        ("Zustand", "Per-user shopping cart state", PINK),
        ("Recharts", "Earnings & dashboard analytics", ACCENT),
        ("TypeScript 5", "Static typing across the entire codebase", BLUE),
    ]
    cols, x0, y0 = 2, Inches(0.9), Inches(1.6)
    cw, ch, gx, gy = Inches(5.65), Inches(1.1), Inches(0.45), Inches(0.18)
    for i, (name, desc, color) in enumerate(techs):
        r, c = divmod(i, cols)
        x = x0 + c * (cw + gx)
        y = y0 + r * (ch + gy)
        card = add_box(s, x, y, cw, ch)
        add_box(s, x, y, Inches(0.1), ch, color=color, shape_type=MSO_SHAPE.RECTANGLE)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Pt(18)
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = name
        run.font.name = FONT_BOLD
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = color
        p.space_after = Pt(2)
        p2 = tf.add_paragraph()
        run2 = p2.add_run()
        run2.text = desc
        run2.font.name = FONT
        run2.font.size = Pt(11.5)
        run2.font.color.rgb = MUTED


def slide_workflow(prs):
    s = blank_slide(prs)
    header(s, "Key Workflows", number=13)

    def flow(y, title, steps, color):
        add_text_box(s, Inches(0.8), y, Inches(11), Inches(0.4),
                     title, size=16, color=color, bold=True, font=FONT_BOLD)
        n = len(steps)
        gap = Inches(0.34)
        total_w = Inches(11.7) - gap * (n - 1)
        bw = Emu(int(total_w) // n)
        x = Inches(0.8)
        for i, step in enumerate(steps):
            card = add_box(s, x, y + Inches(0.5), bw, Inches(0.95), color=CARD)
            card.line.color.rgb = color
            card.line.width = Pt(1.5)
            set_text(card, step, size=12, color=WHITE)
            if i < n - 1:
                add_arrow(s, Emu(int(x) + int(bw) + Inches(0.02)),
                          y + Inches(0.78), gap - Inches(0.06), Inches(0.36),
                          direction="right")
            x = Emu(int(x) + int(bw) + int(gap))

    flow(Inches(1.55), "Material Order Flow", [
        "Browse\nMaterials", "Add to Cart\n(Zustand)", "Checkout", "POST\n/api/orders",
        "MongoDB\norders", "My Activity\nTracking"], ORANGE)
    flow(Inches(3.55), "Service Booking Flow", [
        "Browse\nWorkers", "View Details\n& Pricing", "Book\nService", "POST\n/api/bookings",
        "Provider\nConfirms", "Earnings\nDashboard"], GREEN)
    flow(Inches(5.55), "Authentication Flow", [
        "Register /\nLogin", "bcrypt\nVerify", "JWT Signed\n(7 days)", "HTTP-only\nCookie",
        "Middleware\nGuard", "Role-based\nDashboard"], BLUE)


def slide_screenshot_single(prs, number, title, subtitle, img_name):
    img = shot(img_name)
    if not img:
        return
    s = blank_slide(prs)
    header(s, title, subtitle, number=number)
    iw = Inches(9.2)
    img_obj = Image.open(img)
    ratio = img_obj.size[1] / img_obj.size[0]
    x = Emu(int((SLIDE_W - iw) / 2))
    add_screenshot(s, img, x, Inches(1.55), iw)


def slide_screenshot_pair(prs, number, title, left, right, captions):
    li, ri = shot(left), shot(right)
    if not (li and ri):
        return
    s = blank_slide(prs)
    header(s, title, number=number)
    w = Inches(5.9)
    add_screenshot(s, li, Inches(0.65), Inches(1.7), w, caption=captions[0])
    add_screenshot(s, ri, Inches(6.85), Inches(1.7), w, caption=captions[1])


def slide_testing(prs):
    s = blank_slide(prs)
    header(s, "Testing & Results", number=18)
    styled_table(
        s, Inches(0.8), Inches(1.55), Inches(8.2), Inches(4.9),
        ["Test Case", "Scenario", "Result"],
        [
            ["TC01", "Login with valid / invalid credentials", "Pass"],
            ["TC04", "Add to cart — logged-in user", "Pass"],
            ["TC05", "Cart access as guest → redirect login", "Pass"],
            ["TC06", "Worker booking creation", "Pass"],
            ["TC07", "Material order checkout & total", "Pass"],
            ["TC09", "Provider booking status & earnings", "Pass"],
            ["TC10", "Admin blocks user → login denied", "Pass"],
            ["TC12", "Production build (npm run build)", "Pass"],
        ],
        col_widths=[1.2, 5.4, 1.2],
        header_size=13, body_size=12,
    )
    bullet_card(s, Inches(9.3), Inches(1.55), Inches(3.3), Inches(4.9),
                "Levels Covered", [
            "Unit testing",
            "Integration testing",
            "User acceptance",
            "Output testing",
            "Validation testing",
            "Security checks",
        ], body_size=13)


def slide_conclusion(prs):
    s = blank_slide(prs)
    header(s, "Conclusion", number=19)
    bullet_card(s, Inches(0.7), Inches(1.55), Inches(6.4), Inches(4.9),
                "Achievements", [
            "Unified marketplace: services + materials, one platform",
            "Five-role system with secure JWT access control",
            "Cart, bookings, orders, reviews fully working",
            "Admin / Provider / Vendor dashboards delivered",
            "All 12 test cases passed; clean production build",
        ], title_color=GREEN, body_size=14)
    bullet_card(s, Inches(7.4), Inches(1.55), Inches(5.2), Inches(4.9),
                "Future Enhancements", [
            "Razorpay / UPI payment gateway",
            "Email & SMS notifications",
            "Mobile app (React Native)",
            "GPS worker tracking",
            "AI-based recommendations",
                "Cloud deployment with CI/CD",
        ], title_color=ACCENT, body_size=14)


def slide_thanks(prs):
    s = blank_slide(prs)
    add_box(s, Inches(0), Inches(0), SLIDE_W, Inches(0.18), color=ACCENT,
            shape_type=MSO_SHAPE.RECTANGLE)
    add_box(s, Inches(0), Inches(7.32), SLIDE_W, Inches(0.18), color=ACCENT,
            shape_type=MSO_SHAPE.RECTANGLE)
    add_text_box(s, Inches(1.2), Inches(2.1), Inches(10.93), Inches(1.1),
                 "Thank You", size=60, color=ACCENT, bold=True,
                 align=PP_ALIGN.CENTER, font=FONT_BOLD)
    add_text_box(s, Inches(1.2), Inches(3.4), Inches(10.93), Inches(0.5),
                 "Questions & Discussion", size=20, color=WHITE,
                 align=PP_ALIGN.CENTER)
    card = add_box(s, Inches(3.92), Inches(4.4), Inches(5.5), Inches(1.7))
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (txt, size, color, bold) in enumerate([
        ("HARIPRASAD T B", 17, WHITE, True),
        ("232MCAN0168  |  23MCA5168", 12, MUTED, False),
        ("Guide: Dr. S. Meera — Dept. of Computer Science (AI)", 12, WHITE, False),
        ("Bharathiar University CDOE  |  2023–2025", 12, ACCENT, False),
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = txt
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


# ---------------------------------------------------------------- main

def main():
    crop_screenshots()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_title(prs)            # 1
    slide_agenda(prs)           # 2
    slide_intro(prs)            # 3
    slide_problem(prs)          # 4
    slide_objectives(prs)       # 5
    slide_existing_proposed(prs)  # 6
    slide_architecture(prs)     # 7
    slide_roles(prs)            # 8
    slide_modules(prs)          # 9
    slide_er(prs)               # 10
    slide_collections(prs)      # 11
    slide_tech(prs)             # 12
    slide_workflow(prs)         # 13
    slide_screenshot_single(prs, 14, "Walkthrough — Home Page",
                            "Landing page with search and service categories", "home")
    slide_screenshot_single(prs, 15, "Walkthrough — Services",
                            "Skilled workers with filters, ratings and booking", "workers")
    slide_screenshot_pair(prs, 16, "Walkthrough — Contractors & Login",
                          "contractors", "login",
                          ["Contractor listings", "Login / Registration"])
    slide_screenshot_single(prs, 17, "Walkthrough — Materials Marketplace",
                            "Product catalog with categories, pricing and cart", "materials")
    slide_testing(prs)          # 18
    slide_conclusion(prs)       # 19
    slide_thanks(prs)           # 20

    prs.save(OUTPUT)
    print(f"Created: {OUTPUT} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
